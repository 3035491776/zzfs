# -*- coding: utf-8 -*-
"""Excel 导入导出服务"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def generate_book_template():
    """生成图书导入模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = '图书导入模板'

    # 表头
    headers = ['书名', '作者', 'ISBN', '出版社', '出版年份', '库存', '分类名称', '馆藏位置', '简介']
    header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='409EFF', end_color='409EFF', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # 示例数据
    example = [
        'Python编程入门', '张三', '978-7-111-00001-0', '人民邮电出版社', 2024, 10, '计算机科学', 'A区1排', '一本Python入门书籍'
    ]
    for col, value in enumerate(example, 1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.font = Font(name='微软雅黑', size=11)

    # 列宽
    widths = [30, 15, 25, 20, 12, 8, 15, 15, 40]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # 保存到 bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def parse_book_excel(file_stream):
    """解析图书导入Excel

    Args:
        file_stream: 上传文件流

    Returns:
        (success_count, fail_list): 成功数和失败记录列表
    """
    from openpyxl import load_workbook
    from models.book import Book
    from models.category import Category
    from extensions import db

    wb = load_workbook(file_stream)
    ws = wb.active

    success_count = 0
    fail_list = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row[0]:  # 空行跳过
            continue

        try:
            title = str(row[0]).strip() if row[0] else ''
            author = str(row[1]).strip() if row[1] else ''
            isbn = str(row[2]).strip() if row[2] else ''
            publisher = str(row[3]).strip() if row[3] else ''
            publish_year = int(row[4]) if row[4] else None
            stock = int(row[5]) if row[5] else 0
            cat_name = str(row[6]).strip() if row[6] else ''
            location = str(row[7]).strip() if row[7] else ''
            description = str(row[8]).strip() if row[8] else ''

            if not all([title, author, isbn, cat_name]):
                fail_list.append({'row': row_idx, 'reason': '必填字段缺失'})
                continue

            # 校验ISBN唯一
            if Book.query.filter_by(isbn=isbn).first():
                fail_list.append({'row': row_idx, 'reason': f'ISBN "{isbn}" 已存在'})
                continue

            # 查找分类
            cat = Category.query.filter_by(name=cat_name, is_deleted=False).first()
            if not cat:
                fail_list.append({'row': row_idx, 'reason': f'分类 "{cat_name}" 不存在'})
                continue

            book = Book(
                title=title, author=author, isbn=isbn,
                publisher=publisher, publish_year=publish_year,
                stock=stock, category_id=cat.id,
                location=location, description=description,
            )
            db.session.add(book)
            success_count += 1

        except Exception as e:
            fail_list.append({'row': row_idx, 'reason': str(e)})

    db.session.commit()
    return success_count, fail_list


def export_books_excel():
    """导出图书数据为 Excel bytes"""
    from models.book import Book

    wb = Workbook()
    ws = wb.active
    ws.title = '图书数据'

    # 表头
    headers = ['ID', '书名', '作者', 'ISBN', '出版社', '出版年份', '库存', '分类', '馆藏位置', '入库时间']
    header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='409EFF', end_color='409EFF', fill_type='solid')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 数据
    books = Book.query.filter_by(is_deleted=False).order_by(Book.id.desc()).all()
    for row_idx, book in enumerate(books, 2):
        ws.cell(row=row_idx, column=1, value=book.id)
        ws.cell(row=row_idx, column=2, value=book.title)
        ws.cell(row=row_idx, column=3, value=book.author)
        ws.cell(row=row_idx, column=4, value=book.isbn)
        ws.cell(row=row_idx, column=5, value=book.publisher or '')
        ws.cell(row=row_idx, column=6, value=book.publish_year or '')
        ws.cell(row=row_idx, column=7, value=book.stock)
        ws.cell(row=row_idx, column=8, value=book.category.name if book.category else '')
        ws.cell(row=row_idx, column=9, value=book.location or '')
        ws.cell(row=row_idx, column=10, value=book.create_time.strftime('%Y-%m-%d %H:%M') if book.create_time else '')

    # 列宽
    widths = [8, 30, 15, 25, 20, 12, 8, 15, 15, 20]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_borrows_excel():
    """导出借阅记录为 Excel bytes"""
    from models.borrow import Borrow

    wb = Workbook()
    ws = wb.active
    ws.title = '借阅记录'

    headers = ['ID', '借阅人', '图书', 'ISBN', '状态', '借阅时间', '到期时间', '归还时间', '审核人']
    header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='67C23A', end_color='67C23A', fill_type='solid')

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    borrows = Borrow.query.order_by(Borrow.id.desc()).limit(500).all()
    for row_idx, b in enumerate(borrows, 2):
        ws.cell(row=row_idx, column=1, value=b.id)
        ws.cell(row=row_idx, column=2, value=b.student.real_name if b.student else '')
        ws.cell(row=row_idx, column=3, value=b.book.title if b.book else '')
        ws.cell(row=row_idx, column=4, value=b.book.isbn if b.book else '')
        ws.cell(row=row_idx, column=5, value=b.status_text())
        ws.cell(row=row_idx, column=6, value=b.borrow_time.strftime('%Y-%m-%d %H:%M') if b.borrow_time else '')
        ws.cell(row=row_idx, column=7, value=b.due_time.strftime('%Y-%m-%d') if b.due_time else '')
        ws.cell(row=row_idx, column=8, value=b.return_time.strftime('%Y-%m-%d %H:%M') if b.return_time else '')
        ws.cell(row=row_idx, column=9, value=b.review_admin.real_name if b.review_admin else '')

    widths = [8, 15, 30, 25, 10, 20, 12, 20, 15]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
